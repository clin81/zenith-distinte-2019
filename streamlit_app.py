import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook
from streamlit_gsheets import GSheetsConnection

# ============================================================
# CONFIGURAZIONE PAGINA
# ============================================================

st.set_page_config(
    page_title="Zenith Prato - Gestione Distinte",
    layout="wide"
)

TEMPLATE_FILE = "distinta_vuota.xlsx"

# Nome del foglio Google Sheets da utilizzare
WORKSHEET = "Database_tesserati2019"

# Connessione a Google Sheets
conn = st.connection(
    "gsheets",
    type=GSheetsConnection
)


# ============================================================
# CARICAMENTO DATABASE
# ============================================================

@st.cache_data(ttl=600)
def carica_db_ottimizzato():

    try:

        # ====================================================
        # LEGGE SEMPRE DAL FOGLIO:
        # Database_tesserati2019
        # ====================================================

        df = conn.read(
            worksheet=WORKSHEET,
            ttl="10m"
        )

        # Se il foglio è vuoto
        if df is None or df.empty:

            return pd.DataFrame(
                columns=[
                    "Nominativo",
                    "Tipo",
                    "Ruolo",
                    "Maglia",
                    "GG",
                    "MM",
                    "AA",
                    "FIGC",
                    "Capitano",
                    "Portiere"
                ]
            )

        # ====================================================
        # PULIZIA NOMI COLONNE
        # ====================================================

        df.columns = [
            str(c).strip()
            for c in df.columns
        ]

        # ====================================================
        # GARANZIA COLONNE PRESENTI
        # ====================================================

        col_richieste = [
            "Nominativo",
            "Tipo",
            "Ruolo",
            "Maglia",
            "GG",
            "MM",
            "AA",
            "FIGC",
            "Capitano",
            "Portiere"
        ]

        for col in col_richieste:

            if col not in df.columns:

                if col in ["Capitano", "Portiere"]:
                    df[col] = False
                else:
                    df[col] = ""

        # ====================================================
        # FORMATTAZIONE DATI
        # ====================================================

        df["Tipo"] = (
            df["Tipo"]
            .fillna("")
            .astype(str)
            .str.strip()
        )

        df["Nominativo"] = (
            df["Nominativo"]
            .fillna("")
            .astype(str)
            .str.strip()
        )

        # Numero maglia
        df["Maglia"] = pd.to_numeric(
            df["Maglia"],
            errors="coerce"
        )

        # Capitano / Portiere
        for c in ["Capitano", "Portiere"]:

            df[c] = (
                pd.to_numeric(
                    df[c],
                    errors="coerce"
                )
                .fillna(0)
                .astype(bool)
            )

        # ====================================================
        # ORDINAMENTO
        # ====================================================

        df = df.sort_values(
            by=["Tipo", "Nominativo"],
            ascending=[False, True]
        )

        return df

    except Exception as e:

        st.error(
            f"Errore nel caricamento: {e}"
        )

        return pd.DataFrame()


# ============================================================
# SALVATAGGIO DATABASE
# ============================================================

def salva_db(df):

    try:

        # ====================================================
        # SALVA SEMPRE NEL FOGLIO:
        # Database_tesserati2019
        # ====================================================

        conn.update(
            worksheet=WORKSHEET,
            data=df
        )

        # Svuota la cache
        # così al successivo caricamento
        # vengono letti i dati aggiornati

        st.cache_data.clear()

        return True

    except Exception as e:

        st.error(
            f"Errore nel salvataggio: {e}"
        )

        return False


# ============================================================
# FUNZIONE SCRITTURA EXCEL
# ============================================================

def safe_write(ws, cell_coord, value):

    from openpyxl.cell.cell import MergedCell

    cell = ws[cell_coord]

    if isinstance(cell, MergedCell):

        for m_range in ws.merged_cells.ranges:

            if cell_coord in m_range:

                ws.cell(
                    row=m_range.min_row,
                    column=m_range.min_col
                ).value = value

                return

    else:

        cell.value = value


# ============================================================
# COMPILAZIONE TEMPLATE EXCEL
# ============================================================

def compila_template(p_df, s_df, info):

    wb = load_workbook(
        TEMPLATE_FILE
    )

    ws = wb.active

    # ========================================================
    # INTESTAZIONE GARA
    # ========================================================

    safe_write(
        ws,
        "G7",
        f"Zenith Prato Vs {info['avversario']}"
    )

    safe_write(
        ws,
        "G8",
        f"Data: {info['data']} - Ora: {info['ora']}"
    )

    safe_write(
        ws,
        "G9",
        info["campo"]
    )

    # ========================================================
    # ORDINAMENTO GIOCATORI PER NUMERO MAGLIA
    # ========================================================

    p_df = p_df.sort_values(
        by="Maglia",
        ascending=True
    )

    # ========================================================
    # GIOCATORI
    # ========================================================

    for i, (_, row) in enumerate(
        p_df.iterrows()
    ):

        r = 12 + i

        # Massimo riga 38
        if r > 38:
            break

        nome = f"{row.get('Nominativo', '')}"

        # Capitano
        if row.get("Capitano"):
            nome += " (C)"

        # Portiere
        if row.get("Portiere"):
            nome += " (P)"

        safe_write(
            ws,
            f"C{r}",
            row.get("Maglia", "")
        )

        safe_write(
            ws,
            f"D{r}",
            row.get("GG", "")
        )

        safe_write(
            ws,
            f"E{r}",
            row.get("MM", "")
        )

        safe_write(
            ws,
            f"F{r}",
            row.get("AA", "")
        )

        safe_write(
            ws,
            f"G{r}",
            nome
        )

        safe_write(
            ws,
            f"I{r}",
            row.get("FIGC", "")
        )

    # ========================================================
    # STAFF
    # ========================================================

    for i, (_, row) in enumerate(
        s_df.iterrows()
    ):

        r = 39 + i

        safe_write(
            ws,
            f"C{r}",
            row.get("Ruolo", "")
        )

        safe_write(
            ws,
            f"D{r}",
            row.get("Nominativo", "")
        )

        safe_write(
            ws,
            f"I{r}",
            row.get("FIGC", "")
        )

    # ========================================================
    # CREA FILE IN MEMORIA
    # ========================================================

    out = BytesIO()

    wb.save(out)

    return out.getvalue()


# ============================================================
# INTERFACCIA
# ============================================================

st.title(
    "⚽ Zenith Prato - Sistema Distinte"
)

t1, t2 = st.tabs(
    [
        "📋 Genera Distinta",
        "⚙️ Gestione Anagrafica"
    ]
)


# ============================================================
# TAB 2 - GESTIONE ANAGRAFICA
# ============================================================

with t2:

    st.header(
        "Anagrafica"
    )

    st.caption(
        f"Database: {WORKSHEET}"
    )

    # Carica dati dal foglio Database_tesserati2019
    df = carica_db_ottimizzato()

    if not df.empty:

        # ====================================================
        # GIOCATORI
        # ====================================================

        giocatori_df = df[
            df["Tipo"].str.contains(
                "giocatore",
                case=False,
                na=False
            )
        ]

        giocatori_nomi = (
            giocatori_df["Nominativo"]
            .tolist()
        )

        # ====================================================
        # RUOLI SPECIALI
        # ====================================================

        st.subheader(
            "🏆 Ruoli Speciali"
        )

        c1, c2 = st.columns(2)

        # ----------------------------------------------------
        # CAPITANO
        # ----------------------------------------------------

        with c1:

            cap_lista = df[
                df["Capitano"] == True
            ]["Nominativo"].tolist()

            idx_cap = 0

            if (
                cap_lista
                and cap_lista[0] in giocatori_nomi
            ):

                idx_cap = (
                    giocatori_nomi.index(
                        cap_lista[0]
                    ) + 1
                )

            capitano = st.selectbox(
                "Seleziona Capitano",
                ["Nessuno"] + giocatori_nomi,
                index=idx_cap
            )

        # ----------------------------------------------------
        # PORTIERE
        # ----------------------------------------------------

        with c2:

            por_lista = df[
                df["Portiere"] == True
            ]["Nominativo"].tolist()

            idx_por = 0

            if (
                por_lista
                and por_lista[0] in giocatori_nomi
            ):

                idx_por = (
                    giocatori_nomi.index(
                        por_lista[0]
                    ) + 1
                )

            portiere = st.selectbox(
                "Seleziona Portiere",
                ["Nessuno"] + giocatori_nomi,
                index=idx_por
            )

        # ====================================================
        # TABELLA MODIFICABILE
        # ====================================================

        st.subheader(
            "📋 Anagrafica Tesserati"
        )

        col_vis = [
            "Nominativo",
            "Tipo",
            "Ruolo",
            "Maglia",
            "GG",
            "MM",
            "AA",
            "FIGC"
        ]

        df_edit = st.data_editor(
            df[col_vis],
            num_rows="dynamic",
            use_container_width=True,
            hide_index=True
        )

        # ====================================================
        # SALVATAGGIO
        # ====================================================

        if st.button(
            "💾 Salva e Aggiorna Cache"
        ):

            # Imposta Capitano
            df_edit["Capitano"] = (
                df_edit["Nominativo"]
                == capitano
            )

            # Imposta Portiere
            df_edit["Portiere"] = (
                df_edit["Nominativo"]
                == portiere
            )

            # Salva nel foglio Database_tesserati2019
            if salva_db(df_edit):

                st.success(
                    "Dati salvati correttamente nel foglio "
                    f"'{WORKSHEET}'!"
                )

                st.rerun()


# ============================================================
# TAB 1 - GENERA DISTINTA
# ============================================================

with t1:

    st.header(
        "Gara"
    )

    c1, c2 = st.columns(2)

    # ========================================================
    # DATI GARA
    # ========================================================

    with c1:

        avv = st.text_input(
            "Avversario",
            "SQUADRA OSPITE"
        )

        cmp = st.text_input(
            "Campo",
            "Chiavacci"
        )

    with c2:

        dat = st.text_input(
            "Data",
            "15/04/2026"
        )

        ora = st.text_input(
            "Ora",
            "10:30"
        )

    # ========================================================
    # CARICAMENTO DATABASE
    # ========================================================

    df_gara = carica_db_ottimizzato()

    if not df_gara.empty:

        # ====================================================
        # SEPARA GIOCATORI E STAFF
        # ====================================================

        gioc = df_gara[
            df_gara["Tipo"].str.contains(
                "giocatore",
                case=False,
                na=False
            )
        ]

        staf = df_gara[
            df_gara["Tipo"].str.contains(
                "staff",
                case=False,
                na=False
            )
        ]

        col_s1, col_s2 = st.columns(2)

        # ====================================================
        # SELEZIONE GIOCATORI
        # ====================================================

        with col_s1:

            sel_p = st.multiselect(
                "Giocatori",
                gioc["Nominativo"].tolist()
            )

        # ====================================================
        # SELEZIONE STAFF
        # ====================================================

        with col_s2:

            sel_s = st.multiselect(
                "Staff",
                staf["Nominativo"].tolist()
            )

        # ====================================================
        # GENERAZIONE EXCEL
        # ====================================================

        if st.button(
            "🚀 Genera Excel"
        ):

            if sel_p:

                xlsx = compila_template(

                    gioc[
                        gioc["Nominativo"].isin(sel_p)
                    ],

                    staf[
                        staf["Nominativo"].isin(sel_s)
                    ],

                    {
                        "avversario": avv,
                        "campo": cmp,
                        "data": dat,
                        "ora": ora
                    }
                )

                st.download_button(

                    "📥 Scarica",

                    xlsx,

                    f"Distinta_{avv}.xlsx",

                    mime=(
                        "application/vnd.openxmlformats-"
                        "officedocument.spreadsheetml.sheet"
                    )
                )

            else:

                st.warning(
                    "Seleziona almeno un giocatore."
                )
